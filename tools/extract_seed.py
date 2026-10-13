import openpyxl, json, datetime, sys, warnings
warnings.filterwarnings('ignore')
src=sys.argv[1]; out=sys.argv[2]
wb=openpyxl.load_workbook(src)
def raw(v):
    if v is None: return ''
    if isinstance(v,(datetime.datetime,datetime.date)): return v.strftime('%Y-%m-%d')
    if isinstance(v,str):
        return v[1:].lstrip('+') if v.startswith('=') else v
    return v
def lsrow(d): return 5+d if d<=15 else 7+d
LS_IN=['B','C','E','G','H','K','L','M','O','P','Q','R','S','T','U','V','W','X','Y','Z']
MO_IN=['E','G','H','K','O','P','Q','R','V']
ws=wb['LS_RPRO']
seed={'meta':{'store':'TOBYS MEGAMALL','branch':'MEG','monthStart':ws['A2'].value.strftime('%Y-%m-%d'),
  'target':wb['Seasonality Template']['G15'].value,'gmTarget':wb['Seasonality Template']['M15'].value or 0,
  'runningDays':None,'recipient':'FranchiseDev@tobys.com'}}
seed['ls']={str(d):{c:raw(ws[f'{c}{lsrow(d)}'].value) for c in LS_IN} for d in range(1,32)}
for s,k in [('MANUAL','manual'),('OUTSALE','outsale')]:
    w=wb[s]; seed[k]={str(d):{c:raw(w[f'{c}{lsrow(d)}'].value) for c in MO_IN} for d in range(1,32)}
# MSI / MRet line items (inputs only)
from openpyxl.utils import get_column_letter as L
for s,k in [('MSI','msi'),('MRet','mret')]:
    w=wb[s]; rows=[]
    for r in range(9,303):
        rec={}
        for c in range(1,34):
            col=L(c); v=w.cell(r,c).value
            if col in ('O','Q','R') or (s=='MRet' and col in('A','Z')): continue
            if v is not None and not (isinstance(v,str) and v.startswith('=')): rec[col]=raw(v)
        if rec: rows.append(rec)
    seed[k]=rows
st=wb['Seasonality Template']
S_IN=['G','O','P','T','U','X','AE','AI','AL','AS','AV','BC','BD','BG','BH','BJ','BK','BQ','BR']
seed['seas']={str(d):{c:raw(st[f'{c}{24+d}'].value) for c in S_IN} for d in range(1,32)}
sd=wb['SAME DAY COMPARISON']; seed['same']={str(d):raw(sd[f'B{6+d}'].value) for d in range(1,32)}
eo=wb['EOD MONITORING']; seed['eod']={'start':raw(eo['A3'].value),'rows':{str(d):{'B':raw(eo[f'B{2+d}'].value),'C':raw(eo[f'C{2+d}'].value)} for d in range(1,32)}}
y=wb['YTD-SALES']
seed['ytd']={'months':{str(i):{c:raw(y[f'{c}{6+i}'].value) for c in ['B','C','F','H','I','L','M','O','P']} for i in range(12)},
  'shoes':{str(d):{c:raw(y[f'{c}{2+d}'].value) for c in ['S','T','U','V','W','X','Y','Z','AA','AB','AC','AG']} for d in range(1,32)},
  'labels':{'AF2':raw(y['AF2'].value),'AG2':raw(y['AG2'].value),'AH2':raw(y['AH2'].value),'AJ3':raw(y['AJ3'].value),'AJ4':raw(y['AJ4'].value),'A2':raw(y['A2'].value)}}
yr=wb['YTD RETAIL DRIVERS']
seed['ytdrd']={'months':{str(i):{c:raw(yr[f'{c}{6+i}'].value) for c in ['B','C','E','F','N','O','Q','R','T','U']} for i in range(12)},'A2':raw(yr['A2'].value)}
dp=wb['DEPOSIT SUMMARY']
seed['dep']={'B2':raw(dp['B2'].value),'rows':{str(d):{c:raw(dp[f'{c}{7+d}'].value) for c in ['D','E','F','H','I','J','K','L','M']} for d in range(1,32)}}
sg=wb['#OF SG RECEIPT COPY']
seed['sg']={'B2':raw(sg['B2'].value),'B3':raw(sg['B3'].value),'rows':{str(d):{c:raw(sg[f'{c}{5+d}'].value) for c in ['C','E','F']} for d in range(1,32)}}
json.dump(seed,open(out,'w'),indent=0,default=str)
print('ok', len(seed['msi']), len(seed['mret']))
